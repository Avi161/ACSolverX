"""One sheet: each original, its Aut-minimal form, and what greedy paid.

WHAT THE NAMES MEAN
-------------------
``ac19_<n>``   the n-th ``Aut(F2)`` orbit of ``data/AC19_extended.txt``,
               zero-based, orbits ordered by their first dataset line. NOT a
               line number: ``ac19_50892`` is the 50,892nd orbit and its one
               member is line 90,721. Only ``ac19_0`` coincides with a line.
``ac19x_<n>``  line ``n`` of ``data/AC19_extended.txt``, zero-based.

One row per original, 40 of them over 28 orbits. Columns:

    autmin, autmin_r1, autmin_r2   the orbit and its Aut(F2)-minimal pair,
                                   which greedy failed at 10,000,000 nodes
    original, orig_r1, orig_r2     the raw dataset row it canonicalizes from
    nodes_explored, path_length    greedy on the ORIGINAL, cap 64
    autmin_path_length             the AUT-MIN pair's path length, obtained by
                                   transporting that same certificate through
                                   the aut_canon automorphism and Nielsen-
                                   reducing the basis: path_length + a 1-4
                                   move tail, replayed to (x, y)

``greedy`` only. Read off files already in the repo; the CSV twin is what git
diffs, the workbook is what a person opens.

    PYTHONPATH=. python3 -m experiments.search.make_ac19_orig_10m_xlsx
    PYTHONPATH=. python3 -m experiments.search.make_ac19_orig_10m_xlsx --check
"""
from __future__ import annotations

import argparse
import csv
import io
import os

from experiments.search import make_ac19_orig_10m_lists as mk
from experiments.search import transport_ac19_orig as tr
from experiments.search.run_leftovers_1m import read_rows

OUT_DIR = os.path.join(mk.RESULTS_DIR, "ac19_orig_10m")
XLSX = os.path.join(OUT_DIR, "ac19_orig_10m_originals.xlsx")
CSV = os.path.join(OUT_DIR, "ac19_orig_10m_originals.csv")
ARM = "greedy"
JSONL = "ac19_orig_10m_greedy_b10000000_mrl64.jsonl"

COLUMNS = ("autmin", "autmin_r1", "autmin_r2",
           "original", "orig_r1", "orig_r2",
           "nodes_explored", "path_length", "autmin_path_length")
WIDTHS = {"autmin": 12, "autmin_r1": 16, "autmin_r2": 20,
          "original": 14, "orig_r1": 18, "orig_r2": 18,
          "nodes_explored": 15, "path_length": 12, "autmin_path_length": 19}


def build():
    """The 40 rows, ordered by orbit position then dataset line."""
    want = {r["name"]: r for r in mk.build()[ARM]}
    recs = {r["name"]: r for r in read_rows(os.path.join(OUT_DIR, JSONL))}
    moved = {r["name"]: r for r in read_rows(os.path.join(OUT_DIR, tr.OUT[ARM]))}
    if set(recs) != set(want) or set(moved) != set(want):
        raise RuntimeError("records or transported rows do not cover the "
                           f"{len(want)} derived originals")
    rows = []
    for name, row in want.items():
        rec, mv = recs[name], moved[name]
        if mv["orig_moves"] != rec["path_length"]:
            raise RuntimeError(f"{name}: transported row disagrees with the "
                               "record's path_length")
        rows.append({
            "autmin": row["orbit"],
            "autmin_r1": row["rep_r1"], "autmin_r2": row["rep_r2"],
            "original": name, "orig_r1": row["r1"], "orig_r2": row["r2"],
            "nodes_explored": rec["nodes_explored"],
            "path_length": rec["path_length"],
            "autmin_path_length": mv["rep_moves"],
        })
    rows.sort(key=lambda r: (int(r["autmin"].split("_")[1]),
                             int(r["original"].split("_")[1])))
    return rows


def render_csv(rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\n")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def write_xlsx(rows, path=XLSX):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "originals"
    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
    ws.freeze_panes = "A2"
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[col]
    wb.save(path)
    return path


def read_xlsx(path=XLSX):
    """The sheet as a list of dicts, every cell a string."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    grid = list(ws.iter_rows(values_only=True))
    head = [str(c) for c in grid[0]]
    return [{h: ("" if v is None else str(v)) for h, v in zip(head, r)}
            for r in grid[1:]]


def check(rows=None):
    """``[reason]`` for every way the files on disk differ from the derivation."""
    rows = build() if rows is None else rows
    drift = []
    if not os.path.exists(CSV):
        drift.append(f"{os.path.basename(CSV)} missing")
    elif open(CSV).read() != render_csv(rows):
        drift.append(f"{os.path.basename(CSV)} differs from the derivation")
    if not os.path.exists(XLSX):
        drift.append(f"{os.path.basename(XLSX)} missing")
        return drift
    want = [{c: str(r[c]) for c in COLUMNS} for r in rows]
    got = read_xlsx()
    if got != want:
        bad = next(((i, c) for i, (g, w) in enumerate(zip(got, want))
                    for c in COLUMNS if g.get(c) != w[c]), None)
        drift.append(f"{os.path.basename(XLSX)} differs from the derivation"
                     + (f" (first at row {bad[0] + 2}, column {bad[1]})" if bad else ""))
    return drift


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="report drift, write nothing")
    args = ap.parse_args(argv)
    rows = build()
    print(f"  {len(rows)} originals over "
          f"{len({r['autmin'] for r in rows})} aut-min orbits, {ARM} only")
    if args.check:
        drift = check(rows)
        for d in drift:
            print(f"  DRIFT {d}")
        print("  no drift" if not drift else f"  {len(drift)} difference(s)")
        return 1 if drift else 0
    with open(CSV, "w", newline="") as fh:
        fh.write(render_csv(rows))
    write_xlsx(rows)
    for p in (CSV, XLSX):
        print(f"  wrote {os.path.relpath(p)} ({os.path.getsize(p):,} B)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
